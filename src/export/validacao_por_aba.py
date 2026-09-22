"""Uma planilha de conferência por aba de origem. Não valida nada.

A etapa 1 devolve um `fornecedores_master.csv` só, com tudo junto — é o
que o pipeline precisa, e é ilegível para quem mandou a planilha. Quem
prospectou EPI quer ver EPI: os mesmos sites que ele digitou, na mesma
divisão, com a resposta da Receita ao lado.

Então a regra de nomes é a da origem, não a da categoria FNDE:

    data/output/validacao_lista_sites_EPI.xlsx

O prefixo `validacao_` não é enfeite. As abas da prospecção se chamam
EPI, UNIFORME, Utensílios, Equipamentos — os mesmos nomes das categorias
da entrega. Sem prefixo, `EPI.xlsx` de conferência sobrescreveria a
entrega `EPI.xlsx` no meio do prazo.

Um site que veio de duas abas (astrodistribuidora.com está em EPI e em
UNIFORME) sai nas duas planilhas, com a mesma validação. Repetir é o
certo aqui: cada aba tem que fechar sozinha com o que a pessoa mandou.
"""

from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path

from src.core import tabelas
from src.core.log import obter
from src.models import Fornecedor, Status

log = obter(__name__)

SAIDA = Path("data/output")

COLUNAS = [
    "Empresa", "CNPJ", "Site", "UF", "Status", "Motivo", "CNAE principal",
    "CNAEs secundários", "Situação cadastral", "Atacarejo", "Plataforma",
    "Também em",
]

LARGURAS = {
    "Empresa": 34, "CNPJ": 20, "Site": 38, "UF": 6, "Status": 12,
    "Motivo": 46, "CNAE principal": 15, "CNAEs secundários": 24,
    "Situação cadastral": 20, "Atacarejo": 11, "Plataforma": 14,
    "Também em": 28,
}

# Ordem de leitura: primeiro o que serve, depois o que falta resolver,
# por último o que já foi descartado.
ORDEM_STATUS = {Status.APROVADO: 0, Status.PENDENTE: 1, Status.REPROVADO: 2}

# Sem origem gravada não dá para saber de que aba veio -- é o caso dos
# fornecedores que entraram antes deste campo existir. Ficam num arquivo
# só, em vez de sumir.
SEM_ORIGEM = "sem origem registrada"


def agrupar(fornecedores: list[Fornecedor]) -> dict[str, list[Fornecedor]]:
    """{origem: fornecedores}. Quem veio de duas abas entra nas duas."""
    grupos: dict[str, list[Fornecedor]] = defaultdict(list)
    for forn in fornecedores:
        for origem in forn.origem or [SEM_ORIGEM]:
            grupos[origem].append(forn)
    return grupos


def gerar(fornecedores: list[Fornecedor], destino: Path = SAIDA) -> list[Path]:
    """Escreve uma planilha por origem e devolve os caminhos gerados."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    gerados: list[Path] = []

    for origem, desta in sorted(agrupar(fornecedores).items()):
        wb = Workbook()
        ws = wb.active
        ws.title = _nome_aba(origem)

        ws.append(COLUNAS)
        for celula in ws[1]:
            celula.font = Font(bold=True)
        ws.freeze_panes = "A2"

        for forn in sorted(desta, key=_ordem):
            ws.append(_linha(forn, origem))

        for numero, nome in enumerate(COLUNAS, start=1):
            ws.column_dimensions[get_column_letter(numero)].width = LARGURAS[nome]

        caminho = destino / f"validacao_{_arquivo(origem)}.xlsx"
        wb.save(caminho)
        gerados.append(caminho)

        contagem = {s: sum(1 for f in desta if f.status is s) for s in Status}
        log.info("%s: %d sites (%d aprovados, %d pendentes, %d reprovados) -> %s",
                 origem, len(desta), contagem[Status.APROVADO],
                 contagem[Status.PENDENTE], contagem[Status.REPROVADO], caminho)

    return gerados


def _linha(forn: Fornecedor, origem: str) -> list[str]:
    outras = [o for o in forn.origem if o != origem]
    return [
        forn.nome,
        _cnpj(forn.cnpj),
        forn.url_base,
        forn.uf,
        str(forn.status),
        forn.motivo,
        forn.cnae_principal or "",
        " | ".join(forn.cnaes_secundarios),
        forn.situacao_cadastral or "",
        "SIM" if forn.flag_atacarejo else "",
        str(forn.plataforma or ""),
        ", ".join(_aba(o) for o in outras),
    ]


def _ordem(forn: Fornecedor) -> tuple[int, str]:
    return (ORDEM_STATUS.get(forn.status, 9), (forn.nome or forn.dominio).upper())


def _cnpj(digitos: str | None) -> str:
    """14 dígitos -> 00.000.000/0000-00. Fora disso, devolve como veio.

    O master guarda só dígitos porque é assim que a API consulta; quem
    confere lê com máscara, e é a máscara que ele compara com a planilha
    original. CNPJ que não tem 14 dígitos sai cru de propósito: é
    exatamente o que precisa saltar aos olhos.
    """
    d = digitos or ""
    if len(d) != 14 or not d.isdigit():
        return d
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"


def _aba(origem: str) -> str:
    """"lista_sites.xlsx :: EPI" -> "EPI"."""
    return origem.split("::")[-1].strip()


def _nome_aba(origem: str) -> str:
    """Nome de aba do Excel: 31 caracteres e nada de []:*?/\\."""
    limpo = "".join(" " if c in "[]:*?/\\" else c for c in _aba(origem)).strip()
    return (limpo or "VALIDACAO")[:31]


def _arquivo(origem: str) -> str:
    """Origem -> nome de arquivo previsível, sem acento nem espaço.

    Leva o nome do arquivo de origem junto, e não só o da aba: duas
    prospecções diferentes com uma aba "EPI" cada uma gerariam o mesmo
    `validacao_EPI.xlsx`, e a segunda apagaria a primeira.
    """
    from unidecode import unidecode

    limpo = unidecode(origem).replace(".xlsx", "").replace(".csv", "")
    return "_".join(p for p in "".join(
        c if c.isalnum() else " " for c in limpo
    ).split())


def main(destino: Path = SAIDA) -> list[Path]:
    fornecedores = list(tabelas.ler_fornecedores(apenas_aprovados=False).values())
    gerados = gerar(fornecedores, destino)
    print(f"{len(gerados)} planilha(s) de conferência em {destino}")
    for caminho in gerados:
        print(f"  {caminho}")
    return gerados


if __name__ == "__main__":
    p = argparse.ArgumentParser(
        description="Uma planilha de conferência por aba de origem"
    )
    p.add_argument("--saida", type=Path, default=SAIDA)
    main(p.parse_args().saida)
