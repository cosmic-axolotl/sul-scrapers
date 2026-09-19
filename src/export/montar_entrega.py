"""Monta a entrega a partir de data/coletas/. Não raspa nada.

Lê a pasta inteira, junta com itens.csv e fornecedores_master.csv, corta
nos 5 fornecedores por item e UF, e escreve uma planilha por categoria.

Pode rodar com a coleta pela metade — e deve. Rodar isto no primeiro site
coletado, antes de disparar os outros quarenta, é meia hora que economiza
um dia.

Sobre "os 5": cinco registros não são cinco fornecedores. Ordenar e
cortar em cinco entrega, com facilidade, cinco linhas do mesmo domínio,
todas sem preço, enquanto um candidato com preço fica na reserva. A
entrega pede 5 fornecedores DISTINTOS com preço e print — então o corte
primeiro descarta o que não serve, depois garante um registro por
domínio, e só então conta até cinco.
"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from pathlib import Path

from src.core import tabelas
from src.core.log import obter
from src.export.prints import caminho_miniatura, embutir, gerar_miniatura

log = obter(__name__)

# As 15 colunas do template FNDE_output.xlsx, mais PRINT.
COLUNAS = [
    "Categoria", "Código FGV", "Item", "UF", "Data coleta", "Nome Empresa",
    "CNPJ", "FONTE", "PRODUTO PESQUISADO", "Preço PRODUTO", "VALOR DESCONTO",
    "OBS Desconto", "PREÇO FINAL", "VALOR FRETE", "OBS FRETE", "PRINT",
]

ALVO_POR_ITEM = 5
SCORE_MINIMO = 0.6  # abaixo disso o matching nem classifica como "revisar"

RESERVA = Path("data/interim/reserva.csv")
PENDENCIAS = Path("data/interim/pendencias.csv")


# ---------------------------------------------------------------------------
# Seleção
# ---------------------------------------------------------------------------

def problema(candidato: dict) -> str:
    """Por que este registro NÃO pode ir para a entrega. "" = pode.

    São as três coisas que o solicitante confere numa linha: que tem
    preço, que tem prova, e que o produto é mesmo o pedido.
    """
    preco = candidato.get("preco_final") or candidato.get("preco_produto")
    if preco in (None, "") or _numero(preco) <= 0:
        return "sem preço"

    caminho = (candidato.get("caminho_print") or "").strip()
    if not caminho:
        return "sem print"
    if not Path(caminho).exists():
        return "print não está no disco"

    if candidato.get("valor_frete") in (None, "") and not (
        candidato.get("obs_frete") or ""
    ).strip():
        return "frete não resolvido"

    if _numero(candidato.get("score_match")) < SCORE_MINIMO:
        return f"score {_numero(candidato.get('score_match')):.2f} abaixo do mínimo"

    return ""


def escolher_cinco(candidatos: list[dict]) -> tuple[list[dict], list[dict]]:
    """Os 5 que vão para a entrega e os que sobram para a reserva.

    Critério, nessa ordem: descarta o que não serve; um registro por
    domínio; maior score_match; depois menor preço final.

    A reserva vai para data/interim/reserva.csv e fica só no repositório.
    Ela carrega o motivo de cada exclusão -- sem isso, "por que este
    fornecedor não entrou?" só se responde reprocessando tudo.
    """
    validos: list[dict] = []
    reserva: list[dict] = []

    for c in candidatos:
        motivo = problema(c)
        if motivo:
            reserva.append({**c, "motivo_reserva": motivo})
        else:
            validos.append(c)

    ordenados = sorted(
        validos,
        key=lambda c: (-_numero(c.get("score_match")), _preco(c)),
    )

    escolhidos: list[dict] = []
    dominios: set[str] = set()
    for c in ordenados:
        dominio = c.get("dominio", "")
        if len(escolhidos) < ALVO_POR_ITEM and dominio not in dominios:
            dominios.add(dominio)
            escolhidos.append(c)
        else:
            razao = ("já há registro deste fornecedor" if dominio in dominios
                     else f"além dos {ALVO_POR_ITEM} primeiros")
            reserva.append({**c, "motivo_reserva": razao})

    return escolhidos, reserva


# ---------------------------------------------------------------------------
# Leitura das coletas
# ---------------------------------------------------------------------------

def ler_coletas(entrada: Path) -> list[dict]:
    """Todos os .jsonl da pasta, tolerando linha truncada."""
    from src.runners.etapa3_coletar import ler_registros

    registros: list[dict] = []
    for arquivo in sorted(Path(entrada).glob("*.jsonl")):
        lidos, descartadas = ler_registros(arquivo)
        if descartadas:
            log.warning("%s: %d linhas ilegiveis ignoradas", arquivo.name, descartadas)
        registros.extend(lidos)
    return registros


# ---------------------------------------------------------------------------
# Montagem
# ---------------------------------------------------------------------------

def montar(entrada: Path, saida: Path, categorias: list[str] | None = None) -> dict:
    """Lê os .jsonl, escolhe os 5 e gera um .xlsx por categoria."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    itens = tabelas.ler_itens(categorias)
    fornecedores = tabelas.ler_fornecedores(apenas_aprovados=False)
    coletas = ler_coletas(entrada)

    if not coletas:
        log.warning("nenhuma coleta em %s -- a etapa 3 ja rodou?", entrada)

    # (categoria, id_item, uf) -> candidatos
    grupos: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for c in coletas:
        item = itens.get(c.get("id_item", ""))
        if item is None:
            continue  # item de outra categoria, ou removido da lista
        grupos[(item.categoria, item.id_item, c.get("uf", ""))].append(c)

    reserva_geral: list[dict] = []
    pendencias: list[dict] = []
    linhas_por_categoria: dict[str, list[tuple[dict, dict]]] = defaultdict(list)

    for (categoria, id_item, uf), candidatos in sorted(grupos.items()):
        escolhidos, reserva = escolher_cinco(candidatos)
        reserva_geral.extend(reserva)

        if len(escolhidos) < ALVO_POR_ITEM:
            pendencias.append({
                "categoria": categoria,
                "id_item": id_item,
                "uf": uf,
                "fornecedores": len(escolhidos),
                "faltam": ALVO_POR_ITEM - len(escolhidos),
                "candidatos_descartados": len(reserva),
            })

        item = itens[id_item]
        for c in escolhidos:
            linhas_por_categoria[categoria].append((c, _linha(c, item, fornecedores)))

    saida = Path(saida)
    saida.mkdir(parents=True, exist_ok=True)
    gerados: list[str] = []

    for categoria, linhas in sorted(linhas_por_categoria.items()):
        wb = Workbook()
        ws = wb.active
        ws.title = categoria[:31] or "ENTREGA"

        ws.append(COLUNAS)
        for celula in ws[1]:
            celula.font = Font(bold=True)
        ws.freeze_panes = "A2"

        coluna_print = COLUNAS.index("PRINT") + 1

        for numero, (coleta, linha) in enumerate(linhas, start=2):
            ws.append([linha.get(c) for c in COLUNAS])
            _embutir_print(ws, numero, coluna_print, coleta)

        _ajustar_larguras(ws)
        caminho = saida / f"{_arquivo(categoria)}.xlsx"
        wb.save(caminho)
        gerados.append(str(caminho))
        log.info("%s: %d linhas", caminho, len(linhas))

    _gravar_csv(RESERVA, reserva_geral)
    _gravar_csv(PENDENCIAS, pendencias)

    resumo = {
        "linhas": sum(len(v) for v in linhas_por_categoria.values()),
        "planilhas": gerados,
        "reserva": len(reserva_geral),
        "pendencias": len(pendencias),
    }
    print(
        f"[montar] {resumo['linhas']} linhas em {len(gerados)} planilha(s); "
        f"{resumo['reserva']} na reserva; "
        f"{resumo['pendencias']} combinações item/UF ainda sem 5 fornecedores"
    )
    return resumo


def _linha(coleta: dict, item, fornecedores: dict) -> dict:
    forn = fornecedores.get(coleta.get("dominio", ""))
    return {
        "Categoria": item.categoria,
        "Código FGV": item.id_item,
        "Item": item.item_curto or item.descricao[:90],
        "UF": coleta.get("uf", ""),
        "Data coleta": _data(coleta.get("coletado_em")),
        "Nome Empresa": forn.nome if forn else coleta.get("dominio", ""),
        "CNPJ": _cnpj(forn.cnpj if forn else None),
        "FONTE": coleta.get("url_produto", ""),
        "PRODUTO PESQUISADO": coleta.get("titulo_encontrado", ""),
        "Preço PRODUTO": _numero_ou_nada(coleta.get("preco_produto")),
        "VALOR DESCONTO": _numero_ou_nada(coleta.get("valor_desconto")),
        "OBS Desconto": coleta.get("obs_desconto", ""),
        "PREÇO FINAL": _numero_ou_nada(coleta.get("preco_final")),
        "VALOR FRETE": _numero_ou_nada(coleta.get("valor_frete")),
        "OBS FRETE": coleta.get("obs_frete", ""),
        "PRINT": "",  # a imagem entra ancorada; o texto ficaria atrás dela
    }


def _embutir_print(ws, linha: int, coluna: int, coleta: dict) -> None:
    original = Path((coleta.get("caminho_print") or "").strip())
    if not original.name or not original.exists():
        return
    try:
        miniatura = gerar_miniatura(original, caminho_miniatura(original))
        embutir(ws, linha, coluna, miniatura, original)
    except Exception as e:  # uma miniatura ruim não pode derrubar a entrega
        log.warning("print %s nao entrou na planilha (%s: %s)",
                    original, type(e).__name__, e)


def _ajustar_larguras(ws) -> None:
    from openpyxl.utils import get_column_letter

    larguras = {
        "Categoria": 14, "Código FGV": 12, "Item": 45, "UF": 6,
        "Data coleta": 12, "Nome Empresa": 28, "CNPJ": 20, "FONTE": 40,
        "PRODUTO PESQUISADO": 45, "Preço PRODUTO": 14, "VALOR DESCONTO": 14,
        "OBS Desconto": 24, "PREÇO FINAL": 14, "VALOR FRETE": 12,
        "OBS FRETE": 28,
    }
    for numero, nome in enumerate(COLUNAS, start=1):
        if nome in larguras:
            ws.column_dimensions[get_column_letter(numero)].width = larguras[nome]


def _gravar_csv(caminho: Path, linhas: list[dict]) -> None:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    if not linhas:
        caminho.write_text("", encoding="utf-8")
        return
    colunas = list({chave: None for linha in linhas for chave in linha})
    with caminho.open("w", newline="", encoding="utf-8") as f:
        escritor = csv.DictWriter(f, fieldnames=colunas, extrasaction="ignore")
        escritor.writeheader()
        escritor.writerows(linhas)


# ---------------------------------------------------------------------------

def _numero(valor) -> float:
    try:
        return float(str(valor).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def _numero_ou_nada(valor):
    if valor in (None, ""):
        return None
    try:
        return round(float(str(valor).replace(",", ".")), 2)
    except (TypeError, ValueError):
        return None


def _preco(candidato: dict) -> float:
    preco = candidato.get("preco_final") or candidato.get("preco_produto")
    valor = _numero(preco)
    return valor if valor > 0 else float("inf")


def _data(valor) -> str:
    texto = str(valor or "")
    return texto[:10]


def _cnpj(digitos: str | None) -> str:
    d = "".join(c for c in str(digitos or "") if c.isdigit())
    if len(d) != 14:
        return d
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"


def _arquivo(texto: str) -> str:
    from unidecode import unidecode

    limpo = unidecode(str(texto)).upper()
    return "".join(c if c.isalnum() else "_" for c in limpo).strip("_") or "ENTREGA"


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--entrada", type=Path, default=Path("data/coletas"))
    p.add_argument("--saida", type=Path, default=Path("data/output"))
    p.add_argument("--categoria", action="append", dest="categorias")
    a = p.parse_args()
    montar(a.entrada, a.saida, a.categorias)
