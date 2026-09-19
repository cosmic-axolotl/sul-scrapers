"""Miniatura e hyperlink na coluna PRINT.

A conta que decide tudo: 2.160 prints a 400 KB dariam um .xlsx de ~860 MB,
que não abre. Comprimido a ~50 KB e dividido por categoria, a maior
planilha fica em ~60 MB — abre em qualquer máquina.

Por isso o PNG original fica intocado em data/output/prints/ como prova, e
o que entra na planilha é um JPEG reduzido em data/output/miniaturas/.
"""

from __future__ import annotations

from pathlib import Path

from src.core.log import obter

log = obter(__name__)

LARGURA_MINIATURA = 600
QUALIDADE = 70
ALTURA_LINHA = 110

# Excel mede altura de linha em pontos e largura de coluna em caracteres.
# 1 ponto = 4/3 pixel; a largura de 1 caractere fica em ~7 px na fonte
# padrão. As duas conversões existem só para a miniatura caber na célula.
PONTOS_POR_PIXEL = 0.75
PIXELS_POR_CARACTERE = 7.0

MINIATURAS = Path("data/output/miniaturas")


def caminho_miniatura(original: Path) -> Path:
    """Espelha a estrutura de prints/ dentro de miniaturas/."""
    original = Path(original)
    try:
        relativo = original.relative_to("data/output/prints")
    except ValueError:
        relativo = Path(original.name)
    return MINIATURAS / relativo.with_suffix(".jpg")


def gerar_miniatura(origem: Path, destino: Path | None = None) -> Path:
    """Reduz para LARGURA_MINIATURA e salva JPEG.

    Nunca sobrescreve o PNG original — ele é a prova. Se a miniatura já
    existe e é mais nova que o original, não refaz: montar a entrega é
    uma operação que se roda dezenas de vezes.
    """
    from PIL import Image

    origem = Path(origem)
    destino = Path(destino) if destino else caminho_miniatura(origem)

    if not origem.exists():
        raise FileNotFoundError(f"print nao existe: {origem}")

    if destino.exists() and destino.stat().st_mtime >= origem.stat().st_mtime:
        return destino

    destino.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(origem) as imagem:
        imagem = imagem.convert("RGB")
        imagem.thumbnail((LARGURA_MINIATURA, LARGURA_MINIATURA * 4), Image.LANCZOS)
        imagem.save(destino, "JPEG", quality=QUALIDADE, optimize=True)

    return destino


def embutir(ws, linha: int, coluna: int, miniatura: Path, original: Path) -> None:
    """Ancora a miniatura na célula e aponta o hyperlink para o original.

    A imagem é redimensionada para caber em ALTURA_LINHA: uma captura de
    página inteira tem 3.000 px de altura e, sem isso, cobriria quarenta
    linhas da planilha.
    """
    from openpyxl.drawing.image import Image as ImagemXLSX
    from openpyxl.utils import get_column_letter

    miniatura, original = Path(miniatura), Path(original)
    if not miniatura.exists():
        return

    imagem = ImagemXLSX(str(miniatura))
    proporcao = (imagem.width / imagem.height) if imagem.height else 1.0
    imagem.height = ALTURA_LINHA
    imagem.width = max(1, int(ALTURA_LINHA * proporcao))

    celula = f"{get_column_letter(coluna)}{linha}"
    imagem.anchor = celula
    ws.add_image(imagem)

    ws.row_dimensions[linha].height = ALTURA_LINHA * PONTOS_POR_PIXEL
    largura_atual = ws.column_dimensions[get_column_letter(coluna)].width or 0
    ws.column_dimensions[get_column_letter(coluna)].width = max(
        largura_atual, imagem.width / PIXELS_POR_CARACTERE
    )

    # O hyperlink fica na própria célula PRINT, não numa coluna a mais: o
    # template do FNDE tem 15 colunas e a PRINT é a 16ª combinada. Quem
    # quiser o PNG de prova clica na célula.
    alvo = ws.cell(row=linha, column=coluna)
    try:
        alvo.hyperlink = original.resolve().as_uri()
        alvo.style = "Hyperlink"
    except (OSError, ValueError):
        log.debug("nao consegui montar o hyperlink para %s", original)
