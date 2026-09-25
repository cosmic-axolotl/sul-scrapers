"""Os achados da varredura, em dois arquivos: o que o site tem, e o que revisar.

    data/interim/achados_por_site.csv    uma linha por site, só os ACEITOS
    data/interim/achados_revisar.csv     uma linha por achado DUVIDOSO
    data/output/achados_por_site.xlsx    os mesmos dois, para ler
    data/output/achados_revisar.xlsx

A varredura grava um arquivo por site em data/interim/achados/ -- é o que
permite rodar sites em paralelo sem ninguém escrever no arquivo de
ninguém. Para LER o resultado, um arquivo por site é o formato errado:
são 89 arquivos, 71 deles só com cabeçalho.

Os duvidosos moram em arquivo separado porque a primeira versão os
punha numa coluna ao lado dos aceitos, e ficou ilegível: 160 dos 228
achados eram "revisar", a linha do CSV passava de 2.800 caracteres, e a
coluna deles sumia para fora da tela do editor. Separados, cada um vira
uma linha curta com o que a revisão precisa -- o produto que a varredura
achou, a pontuação e o link. Só o nome do item não diz se o achado está
certo.

Nenhuma saída pode morar dentro de achados/: a varredura e a etapa2_plano
leem aquela pasta inteira, e um arquivo consolidado lá dentro seria lido
como se fosse mais um site.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from src.core import tabelas
from src.core.log import obter

log = obter(__name__)

ACHADOS = Path("data/interim/achados")
SAIDA = Path("data/interim/achados_por_site.csv")
SAIDA_REVISAR = Path("data/interim/achados_revisar.csv")
PLANILHA = Path("data/output/achados_por_site.xlsx")
PLANILHA_REVISAR = Path("data/output/achados_revisar.xlsx")

COLUNAS = ["site", "plataforma", "qtd_aceitos", "qtd_a_revisar", "itens_que_possui"]
COLUNAS_ACHADO = ["site", "plataforma", "item", "produto_encontrado", "score",
                  "url_produto"]


def rotulos_dos_itens() -> dict[str, str]:
    """{id_item: "id - nome curto"}, das duas listas de itens.

    Só o ID ("U043") não diz nada a quem lê; só o nome não deixa achar o
    item na lista FNDE.
    """
    rotulos: dict[str, str] = {}
    for caminho in (tabelas.ITENS, Path("data/interim/itens_varejo.csv")):
        if caminho.exists():
            for id_item, item in tabelas.ler_itens(caminho=caminho).items():
                rotulos.setdefault(id_item, f"{id_item} - {item.item_curto}")
    return rotulos


def ler_achados(pasta: Path = ACHADOS) -> tuple[list[str], dict[str, str], list[dict]]:
    """(sites varridos, {site: plataforma}, um registro por achado)."""
    rotulos = rotulos_dos_itens()
    # A plataforma vem do master. Todos os status, não só APROVADO: um
    # site pode ter achados de uma varredura antiga e ter sido reprovado
    # depois -- a linha dele continua aqui.
    fornecedores = (tabelas.ler_fornecedores(apenas_aprovados=False)
                    if tabelas.FORNECEDORES.exists() else {})

    sites, plataformas, registros = [], {}, []
    for arquivo in sorted(pasta.glob("*.csv")):
        site = arquivo.stem
        sites.append(site)
        plataformas[site] = (str(fornecedores[site].plataforma)
                             if site in fornecedores else "")
        with arquivo.open(encoding="utf-8", newline="") as f:
            for linha in csv.DictReader(f):
                id_item = (linha.get("id_item") or "").strip()
                if not id_item:
                    continue
                try:
                    score = round(float(linha.get("score_match") or 0), 3)
                except ValueError:
                    score = None
                registros.append({
                    "site": site,
                    "plataforma": plataformas[site],
                    "aceito": linha.get("classificacao") == "aceito",
                    "id_item": id_item,
                    "item": rotulos.get(id_item, id_item),
                    "produto_encontrado": linha.get("titulo_encontrado") or "",
                    "score": score,
                    "url_produto": linha.get("url_produto") or "",
                })
    return sites, plataformas, registros


def por_site(sites: list[str], plataformas: dict[str, str],
             registros: list[dict]) -> list[dict]:
    """Uma linha por site, com os itens ACEITOS. Duvidosos só na contagem.

    Sites varridos sem nenhum achado entram também, no fim: "varri e não
    achei" é informação, e é diferente de "ainda não varri".
    """
    def itens(site, aceito):
        vistos = {r["id_item"]: r["item"] for r in registros
                  if r["site"] == site and r["aceito"] is aceito}
        return [vistos[i] for i in sorted(vistos, key=lambda i: (len(i), i))]

    linhas = []
    for site in sites:
        aceitos = itens(site, True)
        linhas.append({
            "site": site,
            "plataforma": plataformas.get(site, ""),
            "qtd_aceitos": len(aceitos),
            "qtd_a_revisar": len(itens(site, False)),
            "itens_que_possui": " | ".join(aceitos),
        })
    # Quem tem mais item confirmado primeiro.
    linhas.sort(key=lambda l: (-l["qtd_aceitos"], -l["qtd_a_revisar"], l["site"]))
    return linhas


def _ordenar(registros: list[dict]) -> list[dict]:
    """Por site; dentro dele, do mais provável para o menos provável."""
    return sorted(registros, key=lambda r: (r["site"], -(r["score"] or 0), r["id_item"]))


def gravar_csv(caminho: Path, colunas: list[str], linhas: list[dict]) -> None:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", newline="", encoding="utf-8") as f:
        escritor = csv.DictWriter(f, fieldnames=colunas, extrasaction="ignore")
        escritor.writeheader()
        escritor.writerows(linhas)


def gravar_planilha(caminho: Path,
                    abas: list[tuple[str, list[str], list[int], list[dict]]]) -> None:
    """Cabeçalho fixo, filtro, larguras certas e o link do produto clicável."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for nome, colunas, larguras, linhas in abas:
        ws = wb.create_sheet(nome)
        ws.append(colunas)
        for c in ws[1]:
            c.font = Font(bold=True)
        for linha in linhas:
            ws.append([linha.get(col, "") for col in colunas])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for n, largura in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(n)].width = largura
        for linha in ws.iter_rows(min_row=2):
            for c in linha:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        # O link clicável é o jeito mais rápido de conferir um achado.
        if "url_produto" in colunas:
            letra = get_column_letter(colunas.index("url_produto") + 1)
            for c in ws[letra][1:]:
                if c.value:
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def main(pasta: Path = ACHADOS) -> None:
    sites, plataformas, registros = ler_achados(pasta)
    resumo = por_site(sites, plataformas, registros)
    aceitos = _ordenar([r for r in registros if r["aceito"]])
    revisar = _ordenar([r for r in registros if not r["aceito"]])

    gravar_csv(SAIDA, COLUNAS, resumo)
    gravar_csv(SAIDA_REVISAR, COLUNAS_ACHADO, revisar)

    larguras_achado = [28, 13, 50, 60, 8, 50]
    # Na planilha, um item por linha dentro da célula: a lista cabe na
    # tela em vez de virar uma faixa de 1.500 caracteres.
    resumo_planilha = [{**l, "itens_que_possui": l["itens_que_possui"].replace(" | ", "\n")}
                       for l in resumo]
    gravar_planilha(PLANILHA, [
        ("Por site", COLUNAS, [30, 14, 11, 13, 70], resumo_planilha),
        ("Aceitos", COLUNAS_ACHADO, larguras_achado, aceitos),
    ])
    gravar_planilha(PLANILHA_REVISAR, [
        ("A revisar", COLUNAS_ACHADO, larguras_achado, revisar),
    ])

    com_achado = sum(1 for l in resumo if l["qtd_aceitos"] or l["qtd_a_revisar"])
    print(f"{len(resumo)} sites varridos, {com_achado} com algum achado")
    print(f"  aceitos : {len(aceitos):4} -> {SAIDA}  e  {PLANILHA}")
    print(f"  revisar : {len(revisar):4} -> {SAIDA_REVISAR}  e  {PLANILHA_REVISAR}")


if __name__ == "__main__":
    argparse.ArgumentParser(
        description="Achados da varredura: o que cada site tem, e o que revisar"
    ).parse_args()
    main()
