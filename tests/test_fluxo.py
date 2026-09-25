"""O fluxo de ponta a ponta, sem rede: varredura -> plano -> coleta -> Excel.

O teste que faltava. A dependência circular (a varredura exigia o
plano_coleta.csv que só a varredura produz) só aparece quando alguém
tenta rodar as duas fases em sequência numa base nova — que é
exatamente o que este arquivo faz.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest

from src.core import tabelas
from src.export import montar_entrega
from src.models import Achado, Fornecedor, Plataforma, Status
from src.orquestrador import (
    montar_lotes,
    montar_lotes_varredura,
    pertence_ao_shard,
)
from src.runners import etapa2_plano, etapa2_varredura


@pytest.fixture
def projeto(tmp_path, monkeypatch):
    """Um projeto inteiro em miniatura: 2 itens, 2 fornecedores aprovados."""
    monkeypatch.chdir(tmp_path)
    (tmp_path / "data/interim/achados").mkdir(parents=True)
    (tmp_path / "data/coletas").mkdir(parents=True)

    itens = tmp_path / "data/interim/itens.csv"
    with itens.open("w", newline="", encoding="utf-8") as f:
        escritor = csv.writer(f)
        escritor.writerow(["id_item", "categoria", "grupo_insumo", "item_curto",
                           "descricao", "termos_busca"])
        escritor.writerow(["U001", "UTENSILIOS", "TABUA", "Tábua de corte",
                           "Tabua de corte polietileno 50 cm", "tabua corte 50 cm"])
        escritor.writerow(["294", "EPI", "MASCARA", "Máscara descartável",
                           "Mascara descartavel TNT 100UNIDADE COMERCIALIZADO EM PACOTE",
                           "mascara 100 unidades"])

    tabelas.gravar_fornecedores(
        [
            Fornecedor(nome="Loja A", dominio="loja-a.com.br",
                       url_base="https://loja-a.com.br", uf="PR",
                       cnpj="61340901000117", plataforma=Plataforma.VTEX,
                       status=Status.APROVADO),
            Fornecedor(nome="Loja B", dominio="loja-b.com.br",
                       url_base="https://loja-b.com.br", uf="SC",
                       cnpj="82922212000190", plataforma=Plataforma.VTEX,
                       status=Status.APROVADO),
            Fornecedor(nome="Reprovada", dominio="loja-c.com.br",
                       url_base="https://loja-c.com.br", uf="RS",
                       status=Status.REPROVADO, motivo="varejo puro"),
        ],
        tmp_path / "data/interim/fornecedores_master.csv",
    )
    return tmp_path


# ---------------------------------------------------------------------------
# Etapa 2 — varredura não depende do plano de coleta
# ---------------------------------------------------------------------------

def test_varredura_roda_sem_plano_de_coleta(projeto):
    """A dependência circular: antes, isto era impossível numa base nova."""
    assert not (projeto / "data/interim/plano_coleta.csv").exists()

    lotes = montar_lotes("varredura", categorias=["UTENSILIOS"])

    assert set(lotes) == {"loja-a.com.br", "loja-b.com.br"}
    assert all(ids == ["U001"] for ids in lotes.values())


def test_varredura_nao_inclui_fornecedor_reprovado(projeto):
    lotes = montar_lotes_varredura()
    assert "loja-c.com.br" not in lotes


def test_coleta_sem_plano_diz_o_que_fazer(projeto):
    """Erro que ensina o próximo passo, em vez de KeyError."""
    with pytest.raises(FileNotFoundError, match="varredura"):
        montar_lotes("coleta")


def test_limite_corta_os_itens_nao_os_sites(projeto):
    lotes = montar_lotes("varredura", limite=1)
    assert len(lotes) == 2
    assert all(len(ids) == 1 for ids in lotes.values())


def test_shard_particiona_os_sites_de_forma_estavel():
    dominios = [f"loja{i}.com.br" for i in range(60)]
    fatias = [
        {d for d in dominios if pertence_ao_shard(d, n, 3)} for n in (1, 2, 3)
    ]

    assert set().union(*fatias) == set(dominios)  # ninguém fica de fora
    assert sum(len(f) for f in fatias) == len(dominios)  # ninguém em duas
    assert all(fatias)  # e a divisão não degenera


# ---------------------------------------------------------------------------
# Etapa 2 — plano
# ---------------------------------------------------------------------------

def _varrer(dominio: str, id_item: str, classificacao: str = "aceito") -> None:
    etapa2_varredura.gravar(
        dominio,
        [Achado(
            id_item=id_item,
            dominio=dominio,
            url_produto=f"https://{dominio}/{id_item}",
            titulo_encontrado=f"Produto {id_item}",
            score_match=0.92,
            preco_indicativo=49.9,
            classificacao=classificacao,
        )],
        [id_item],
    )


def test_plano_nasce_dos_achados_da_varredura(projeto):
    _varrer("loja-a.com.br", "U001")
    _varrer("loja-b.com.br", "U001")

    etapa2_plano.gerar_plano()

    with (projeto / "data/interim/plano_coleta.csv").open(encoding="utf-8") as f:
        linhas = {l["dominio"]: l for l in csv.DictReader(f)}

    assert set(linhas) == {"loja-a.com.br", "loja-b.com.br"}
    assert linhas["loja-a.com.br"]["ids_itens"] == "U001"
    assert linhas["loja-a.com.br"]["uf"] == "PR"


def test_achado_em_revisao_fica_fora_do_plano(projeto):
    """Mandar "revisar" para a coleta é o mesmo que não ter revisão."""
    _varrer("loja-a.com.br", "U001", classificacao="revisar")

    etapa2_plano.gerar_plano()

    with (projeto / "data/interim/plano_coleta.csv").open(encoding="utf-8") as f:
        assert list(csv.DictReader(f)) == []

    with (projeto / "data/interim/revisar.csv").open(encoding="utf-8") as f:
        assert [l["id_item"] for l in csv.DictReader(f)] == ["U001"]


def test_achado_em_site_nao_aprovado_fica_fora_do_plano(projeto):
    """Regra inviolável 2, verificada onde importa."""
    _varrer("loja-c.com.br", "U001")

    etapa2_plano.gerar_plano()

    with (projeto / "data/interim/plano_coleta.csv").open(encoding="utf-8") as f:
        assert list(csv.DictReader(f)) == []


def test_depois_do_plano_a_coleta_encontra_o_trabalho(projeto):
    """A outra ponta da dependência circular."""
    _varrer("loja-a.com.br", "U001")
    etapa2_plano.gerar_plano()

    lotes = montar_lotes("coleta", categorias=["UTENSILIOS"])

    assert lotes == {"loja-a.com.br": ["U001"]}


def test_cobertura_conta_fornecedores_distintos(projeto):
    _varrer("loja-a.com.br", "U001")
    _varrer("loja-b.com.br", "U001")

    resumo = etapa2_plano.medir_cobertura()

    assert resumo["UTENSILIOS"]["completos"] == 0  # 2 fontes, a meta são 5
    assert resumo["UTENSILIOS"]["faltando"] == [("U001", 2)]
    assert resumo["EPI"]["faltando"] == [("294", 0)]


# ---------------------------------------------------------------------------
# Montagem
# ---------------------------------------------------------------------------

def _coletar(projeto: Path, dominio: str, id_item: str, uf: str, preco: float) -> None:
    prints = projeto / "data/output/prints/UTENSILIOS"
    prints.mkdir(parents=True, exist_ok=True)
    imagem = prints / f"{id_item}__{dominio}__{uf}.png"
    imagem.write_bytes(_png_minimo())

    registro = {
        "id_item": id_item, "dominio": dominio, "uf": uf,
        "url_produto": f"https://{dominio}/{id_item}",
        "titulo_encontrado": "Tabua de corte 50 cm",
        "preco_produto": preco, "valor_desconto": None, "obs_desconto": "",
        "preco_final": preco, "valor_frete": 20.0, "obs_frete": "tabela",
        "caminho_print": str(imagem), "score_match": 0.92,
        "coletado_em": "2026-09-18T10:00:00", "categoria": "UTENSILIOS",
    }
    arquivo = projeto / "data/coletas" / f"{dominio}.jsonl"
    with arquivo.open("a", encoding="utf-8") as f:
        f.write(json.dumps(registro, ensure_ascii=False) + "\n")


def _png_minimo() -> bytes:
    """1x1 branco, para a miniatura ter o que reduzir."""
    import base64

    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmM"
        "IQAAAABJRU5ErkJggg=="
    )


def test_monta_a_planilha_com_as_16_colunas(projeto):
    from openpyxl import load_workbook

    _coletar(projeto, "loja-a.com.br", "U001", "PR", 49.9)

    resumo = montar_entrega.montar(
        projeto / "data/coletas", projeto / "data/output", ["UTENSILIOS"]
    )

    assert resumo["linhas"] == 1
    planilha = projeto / "data/output/UTENSILIOS.xlsx"
    assert planilha.exists()

    ws = load_workbook(planilha).active
    assert [c.value for c in ws[1]] == montar_entrega.COLUNAS
    linha = {c.value for c in ws[2]}
    assert "U001" in linha
    assert 49.9 in linha


def test_monta_com_a_coleta_pela_metade_e_registra_a_pendencia(projeto):
    """Rodar isto no primeiro site coletado é o ponto da etapa."""
    _coletar(projeto, "loja-a.com.br", "U001", "PR", 49.9)

    resumo = montar_entrega.montar(
        projeto / "data/coletas", projeto / "data/output", ["UTENSILIOS"]
    )

    assert resumo["pendencias"] == 1
    with (projeto / "data/interim/pendencias.csv").open(encoding="utf-8") as f:
        pendencia = next(csv.DictReader(f))
    assert pendencia["faltam"] == "4"


def test_a_entrega_nunca_repete_fornecedor_no_mesmo_item_e_uf(projeto):
    from openpyxl import load_workbook

    for numero in range(3):
        _coletar(projeto, "loja-a.com.br", "U001", "PR", 40.0 + numero)

    montar_entrega.montar(
        projeto / "data/coletas", projeto / "data/output", ["UTENSILIOS"]
    )

    ws = load_workbook(projeto / "data/output/UTENSILIOS.xlsx").active
    assert ws.max_row == 2  # cabeçalho + uma linha só


# ---------------------------------------------------------------------------
# --so-com-cnpj: não garimpar CNPJ no site
# ---------------------------------------------------------------------------

def test_sem_cnpj_fica_pendente_sem_tocar_no_site(cliente_falso):
    """O passo mais caro da etapa 1: 102 empresas da base regional sem
    CNPJ custam mais de uma hora de raspagem."""
    from src.models import Fornecedor, Status
    from src.runners.etapa1_validar import validar

    cliente = cliente_falso()
    forn = Fornecedor(nome="Alfa", dominio="alfa.com.br",
                      url_base="https://alfa.com.br", uf="PR")

    validar(forn, cliente, buscar_cnpj_no_site=False)

    assert forn.status is Status.PENDENTE
    assert "sem CNPJ na planilha" in forn.motivo
    assert cliente.pedidos == []          # nenhuma requisição


def test_com_cnpj_continua_sendo_validado(cliente_falso, tmp_path, monkeypatch):
    from src.models import Fornecedor, Status
    from src.runners.etapa1_validar import validar

    # A consulta lê data/raw/cnpj/ antes da rede: sem isolar o cache, o
    # teste responde com o CNAE real que estiver gravado na máquina.
    monkeypatch.setattr("src.validacao.consultar_cnpj.CACHE", tmp_path / "cnpj")

    cliente = cliente_falso({"minhareceita": {
        "razao_social": "ALFA LTDA", "cnae_fiscal": 4642702,
        "descricao_situacao_cadastral": "ATIVA", "cnaes_secundarios": [],
    }})
    forn = Fornecedor(nome="Alfa", dominio="alfa.com.br",
                      url_base="https://alfa.com.br", uf="PR",
                      cnpj="61340901000117")

    validar(forn, cliente, buscar_cnpj_no_site=False)

    assert forn.status is Status.APROVADO
    assert forn.cnae_principal == "4642702"


def test_por_padrao_ainda_garimpa_no_site(cliente_falso):
    """Quem não pedir o atalho continua com o comportamento de sempre."""
    from src.models import Fornecedor
    from src.runners.etapa1_validar import validar

    cliente = cliente_falso()
    forn = Fornecedor(nome="Alfa", dominio="alfa.com.br",
                      url_base="https://alfa.com.br", uf="PR")

    validar(forn, cliente)

    assert cliente.pedidos          # tentou ler o site
