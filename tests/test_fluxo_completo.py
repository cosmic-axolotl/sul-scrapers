"""O fluxo inteiro contra uma loja de verdade — só que a loja é nossa.

É a prioridade da análise virada em teste: um site, um item, as três UFs,
com Excel e prints. Sobe um servidor HTTP local que fala VTEX, roda
varredura -> plano -> coleta -> entrega, e confere o resultado.

Marcado como `lento` porque abre um navegador de verdade (~10 s):

    pytest -m "not lento"      # o resto da suíte, em segundos
    pytest -m lento            # só este
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest

from src.core import tabelas
from src.export import montar_entrega
from src.models import Fornecedor, Plataforma, Status
from src.orquestrador import montar_lotes, rodar_site
from src.runners import etapa2_plano
from tests.loja_falsa import subir

pytestmark = pytest.mark.lento


def _tem_navegador() -> bool:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return False
    try:
        with sync_playwright() as p:
            navegador = p.chromium.launch(headless=True)
            navegador.close()
        return True
    except Exception:
        return False


sem_navegador = pytest.mark.skipif(
    not _tem_navegador(),
    reason="precisa de playwright + chromium (pip install playwright && "
           "playwright install chromium)",
)

DESCRICAO = (
    "Kit de tabuas profissionais para corte de alimentos, composto por 5 "
    "unidades em cores distintas, em polietileno de alta densidade. "
    "Dimensoes aproximadas de 50 x 30 x 0,8 cm"
)


@pytest.fixture
def loja(tmp_path, monkeypatch):
    """Projeto em miniatura apontando para a loja falsa."""
    servidor, base = subir()
    monkeypatch.chdir(tmp_path)

    Path("data/interim").mkdir(parents=True)
    Path("data/coletas").mkdir(parents=True)

    with open("data/interim/itens.csv", "w", newline="", encoding="utf-8") as f:
        escritor = csv.writer(f)
        escritor.writerow(["id_item", "categoria", "grupo_insumo", "item_curto",
                           "descricao", "termos_busca"])
        escritor.writerow(["U053", "UTENSILIOS", "TABUA DE CORTE",
                           "Kit de tabuas profissionais", DESCRICAO,
                           "tabua corte 5 unidades|tabua corte 50 cm|tabua corte"])

    tabelas.gravar_fornecedores([Fornecedor(
        nome="Loja Falsa Utensilios", dominio="127.0.0.1", url_base=base, uf="PR",
        cnpj="61340901000117", cnae_principal="4641902",
        situacao_cadastral="ATIVA", eh_atacadista=True,
        plataforma=Plataforma.VTEX, status=Status.APROVADO,
    )])

    yield base
    servidor.shutdown()


def _varrer():
    for dominio, ids in montar_lotes("varredura", categorias=["UTENSILIOS"]).items():
        rodar_site(dominio, ids, "varredura")


def _coletar():
    resultados = []
    for dominio, ids in montar_lotes("coleta", categorias=["UTENSILIOS"]).items():
        resultados.append(rodar_site(dominio, ids, "coleta"))
    return resultados


# ---------------------------------------------------------------------------
# Varredura + plano (não precisam de navegador)
# ---------------------------------------------------------------------------

def test_varredura_escolhe_o_kit_e_nao_a_unidade_avulsa(loja):
    """items[0] é a tábua de 30 cm avulsa. O item pede o kit de 5."""
    _varrer()

    with open("data/interim/achados/127.0.0.1.csv", encoding="utf-8") as f:
        achado = next(csv.DictReader(f))

    assert "Kit 5 unidades" in achado["titulo_encontrado"]
    assert achado["sku"] == "5002"
    assert achado["classificacao"] == "aceito"


def test_varredura_pega_o_preco_do_seller_com_estoque(loja):
    """sellers[0] do kit está sem estoque e custa R$ 250,00."""
    _varrer()

    with open("data/interim/achados/127.0.0.1.csv", encoding="utf-8") as f:
        achado = next(csv.DictReader(f))

    assert float(achado["preco_indicativo"]) == 189.9


def test_plano_sai_da_varredura_e_alimenta_a_coleta(loja):
    _varrer()
    etapa2_plano.gerar_plano()

    assert montar_lotes("coleta") == {"127.0.0.1": ["U053"]}


# ---------------------------------------------------------------------------
# Coleta e entrega (precisam de navegador para o print)
# ---------------------------------------------------------------------------

@sem_navegador
def test_coleta_grava_as_tres_ufs_com_frete_diferente(loja):
    _varrer()
    etapa2_plano.gerar_plano()

    [resultado] = _coletar()
    assert resultado["gravados"] == 3
    assert resultado["falhas"] == 0

    linhas = Path("data/coletas/127.0.0.1.jsonl").read_text(
        encoding="utf-8").strip().splitlines()
    por_uf = {json.loads(l)["uf"]: json.loads(l) for l in linhas}

    assert set(por_uf) == {"PR", "SC", "RS"}
    assert por_uf["PR"]["valor_frete"] == 24.9
    assert por_uf["SC"]["valor_frete"] == 31.5
    assert por_uf["RS"]["valor_frete"] == 38.9


@sem_navegador
def test_coleta_separa_preco_riscado_de_preco_pago(loja):
    _varrer()
    etapa2_plano.gerar_plano()
    _coletar()

    registro = json.loads(
        Path("data/coletas/127.0.0.1.jsonl").read_text(encoding="utf-8").splitlines()[0]
    )
    assert registro["preco_produto"] == 229.9
    assert registro["preco_final"] == 189.9
    assert registro["valor_desconto"] == 40.0


@sem_navegador
def test_print_existe_um_por_uf(loja):
    _varrer()
    etapa2_plano.gerar_plano()
    _coletar()

    prints = sorted(Path("data/output/prints").rglob("*.png"))
    assert [p.name for p in prints] == [
        "U053__127.0.0.1__PR.png",
        "U053__127.0.0.1__RS.png",
        "U053__127.0.0.1__SC.png",
    ]
    assert all(p.stat().st_size > 1000 for p in prints)


@sem_navegador
def test_rerun_nao_refaz_o_que_ja_esta_completo(loja):
    """A retomada, medida onde ela importa."""
    _varrer()
    etapa2_plano.gerar_plano()
    _coletar()

    [segunda] = _coletar()
    assert segunda["gravados"] == 0


@sem_navegador
def test_print_apagado_e_refeito_no_rerun(loja):
    """Registro sem prova não é registro pronto."""
    _varrer()
    etapa2_plano.gerar_plano()
    _coletar()

    Path("data/output/prints/UTENSILIOS/U053__127.0.0.1__SC.png").unlink()

    [segunda] = _coletar()
    assert segunda["gravados"] == 1


@sem_navegador
def test_entrega_sai_com_as_16_colunas_e_os_prints_embutidos(loja):
    from openpyxl import load_workbook

    _varrer()
    etapa2_plano.gerar_plano()
    _coletar()

    resumo = montar_entrega.montar(
        Path("data/coletas"), Path("data/output"), ["UTENSILIOS"]
    )
    assert resumo["linhas"] == 3

    ws = load_workbook("data/output/UTENSILIOS.xlsx").active
    assert [c.value for c in ws[1]] == montar_entrega.COLUNAS
    assert ws.max_row == 4  # cabeçalho + 3 UFs
    assert len(ws._images) == 3  # uma miniatura por linha

    ufs = {ws.cell(row=n, column=montar_entrega.COLUNAS.index("UF") + 1).value
           for n in (2, 3, 4)}
    assert ufs == {"PR", "SC", "RS"}
